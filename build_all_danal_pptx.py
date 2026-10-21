import os
import sys
import io
import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

TEMPLATE_PATH = r"C:\Users\USER\Downloads\신용카드 검수요청서(양식).pptx"
DOWNLOADS_DIR = r"C:\Users\USER\Downloads"
APP_DIR = r"C:\App\LuckyPick"

# Dimensions
MAX_W = Emu(9106000)
MAX_H = Emu(5700000)
BASE_TOP = Emu(880000)
BASE_LEFT = Emu(400000)

def delete_slide(prs, index):
    rId = prs.slides._sldIdLst[index].rId
    prs.part.drop_rel(rId)
    del prs.slides._sldIdLst[index]

def setup_header(slide, title_text, top_right_text=None, subtitle_text=None, test_account=False):
    # Remove old body shapes
    shapes_to_remove = []
    title_shape = None
    for shape in slide.shapes:
        if shape.has_text_frame and shape.top < Emu(750000):
            title_shape = shape
        else:
            shapes_to_remove.append(shape)

    for s in shapes_to_remove:
        sp = s._element
        sp.getparent().remove(sp)

    if not title_shape:
        title_shape = slide.shapes.add_textbox(Emu(350000), Emu(120000), Emu(6000000), Emu(650000))

    tf = title_shape.text_frame
    tf.word_wrap = True
    tf.text = title_text
    p0 = tf.paragraphs[0]
    p0.font.bold = True
    p0.font.size = Pt(18)
    p0.font.color.rgb = RGBColor(30, 41, 59)

    if subtitle_text:
        p1 = tf.add_paragraph()
        p1.text = subtitle_text
        p1.font.bold = False
        p1.font.size = Pt(12)
        p1.font.color.rgb = RGBColor(71, 85, 105)

    if top_right_text:
        tb_tr = slide.shapes.add_textbox(Emu(5000000), Emu(100000), Emu(4500000), Emu(650000))
        tf_tr = tb_tr.text_frame
        tf_tr.word_wrap = True
        tf_tr.text = top_right_text
        for p in tf_tr.paragraphs:
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = RGBColor(30, 41, 59)
            p.alignment = PP_ALIGN.RIGHT

    if test_account:
        tb_test = slide.shapes.add_textbox(Emu(6800000), Emu(120000), Emu(2800000), Emu(500000))
        tf_t = tb_test.text_frame
        tf_t.text = "TEST ID : tester@luckypick.kr\nTEST PW : LuckyPick2026!"
        for p in tf_t.paragraphs:
            p.font.size = Pt(11)
            p.font.bold = True
            p.font.color.rgb = RGBColor(220, 38, 38)
            p.alignment = PP_ALIGN.RIGHT

def add_image(slide, img_path, custom_top=None, custom_max_h=None):
    if not os.path.exists(img_path):
        print(f"Warning: Image not found: {img_path}")
        return
    with Image.open(img_path) as im:
        w_px, h_px = im.size

    cur_top = custom_top if custom_top else BASE_TOP
    cur_max_h = custom_max_h if custom_max_h else MAX_H

    scale_w = MAX_W / w_px
    scale_h = cur_max_h / h_px
    scale = min(scale_w, scale_h)

    final_w = int(w_px * scale)
    final_h = int(h_px * scale)

    pos_left = int(BASE_LEFT + (MAX_W - final_w) / 2)
    pos_top = int(cur_top + (cur_max_h - final_h) / 2)

    slide.shapes.add_picture(img_path, pos_left, pos_top, final_w, final_h)

# ===========================================================================
# 1. BUILD COMPLETE 11-SLIDE MASTER PRESENTATION
# ===========================================================================
def build_master_11_slides():
    prs = pptx.Presentation(TEMPLATE_PATH)
    while len(prs.slides) > 11:
        delete_slide(prs, len(prs.slides) - 1)
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    while len(prs.slides) < 11:
        prs.slides.add_slide(blank_layout)

    # Slide 1: Cover
    s0 = prs.slides[0]
    for shape in s0.shapes:
        if shape.has_text_frame:
            text = shape.text.strip()
            tf = shape.text_frame
            if "업체명" in text or "다날" in text or "상호명" in text or shape.name == "텍스트 상자 28":
                tf.text = "럭키픽 (LuckyPick)\n다날 결제서비스 연동 검수요청서"
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].font.size = Pt(28)
                tf.paragraphs[0].font.color.rgb = RGBColor(26, 32, 44)
            elif "URL" in text or "TEST" in text or "작성" in text:
                tf.text = (
                    "서비스 URL: https://luckypick-ec4cf.web.app\n"
                    "상호(법인명): 럭키픽 | 대표자: 이재영\n"
                    "사업자등록번호: 803-05-03449 | 통신판매업: 제 2026-화성새솔-0099 호\n"
                    "고객센터: 070-8019-3205 | 이메일: leejeayoung0713@gmail.com\n"
                    "사업장 소재지: 경기도 화성시 만세구 수노을1로 148, 101동 902호\n"
                    "TEST 계정: ID: tester@luckypick.kr / PW: LuckyPick2026!\n"
                    "신청 서비스: 휴대폰 소액결제(APP), 신용카드 결제(일반/간편)\n"
                    "작성일자: 2026. 10. 21"
                )
                for p in tf.paragraphs:
                    p.font.size = Pt(13)
                    p.font.color.rgb = RGBColor(55, 65, 81)

    # Slide 2 ~ 7: Standard Flow
    std_slides = [
        {
            "index": 1,
            "title": "① 사이트 메인 (URL노출 캡쳐)",
            "subtitle": "서비스 URL (https://luckypick-ec4cf.web.app) 상단 주소창 노출",
            "img": r"C:\App\LuckyPick\danal_captures\01_main_home.png",
            "test_account": False
        },
        {
            "index": 2,
            "title": "② 사이트 메인 하단정보 ※상호명, 사업자번호, 유선번호, 소재지, 대표자명 기재必",
            "subtitle": "상호: 럭키픽 | 대표: 이재영 | 사업자번호: 803-05-03449 | 고객센터: 070-8019-3205",
            "img": r"C:\App\LuckyPick\danal_captures\02_footer_business_info.png",
            "test_account": False
        },
        {
            "index": 3,
            "title": "③ 상품 및 포인트 충전 결제 프로세스 ※ 결제진행상품금액동일/ 금액임의입력불가",
            "subtitle": "정찰제 패키지 선택 (10,000P / 32,000P / 55,000P / 115,000P), 금액 임의 수정 불가",
            "img": r"C:\App\LuckyPick\danal_captures\07_point_charge_section.png",
            "test_account": True
        },
        {
            "index": 4,
            "title": "④ 결제창 연동 (포인트 충전 ₩10,000원 결제 수단 선택)",
            "subtitle": "정찰제 ₩10,000원 결제 대상 > 신용·체크카드 / 휴대폰 소액결제 선택창 호출",
            "img": r"C:\App\LuckyPick\danal_captures\08_point_charge_modal.png",
            "test_account": True
        },
        {
            "index": 5,
            "title": "⑤ 각 카드사 인증창 호출 (다날 신용카드 표준결제창 호출)",
            "subtitle": "다날 신용카드 표준 결제창 호출 > 각 카드사 선택창 (10,000원 결제)",
            "img": r"C:\Users\USER\.gemini\antigravity-ide\brain\d7311755-2ddf-4acd-8b6e-a84b6e5f359a\.user_uploaded\media_1792501582882.png",
            "test_account": False
        },
        {
            "index": 6,
            "title": "⑤ 휴대폰 소액결제 연동창 호출 (SKT, KT, LGU+, 알뜰폰)",
            "subtitle": "다날 휴대폰 소액결제 연동창 호출 > 통신 3사 및 알뜰폰 본인인증 결제창",
            "img": r"C:\Users\USER\.gemini\antigravity-ide\brain\d7311755-2ddf-4acd-8b6e-a84b6e5f359a\.user_uploaded\media_1792501062466.png",
            "test_account": False
        }
    ]

    for cfg in std_slides:
        s = prs.slides[cfg["index"]]
        setup_header(s, cfg["title"], subtitle_text=cfg.get("subtitle"), test_account=cfg.get("test_account", False))
        add_image(s, cfg["img"])

    # Slide 8 (Index 7): 포인트 사용 프로세스 (Exact twin of Danal Sample Page 8)
    s8 = prs.slides[7]
    setup_header(
        s8,
        title_text="포인트 사용 프로세스",
        top_right_text="원하는 서비스 및 수량 선택 후\n결제를 진행하면 충전된 포인트에서 차감됩니다"
    )
    add_image(s8, r"C:\App\LuckyPick\danal_captures\08_point_use_process_v2.png", custom_top=Emu(950000), custom_max_h=Emu(5600000))

    # Slide 9 (Index 8): 이용약관(환불) (Exact twin of Danal Sample Page 9)
    s9 = prs.slides[8]
    setup_header(s9, title_text="이용약관(환불)")
    add_image(s9, r"C:\App\LuckyPick\danal_captures\09_refund_terms_v2.png", custom_top=Emu(950000), custom_max_h=Emu(5600000))

    # Slide 10 (Index 9): 포인트 사용처 (Exact twin of Danal Sample Page 10)
    s10 = prs.slides[9]
    setup_header(s10, title_text="포인트 사용처")
    add_image(s10, r"C:\App\LuckyPick\danal_captures\10_point_scope_v2.png", custom_top=Emu(950000), custom_max_h=Emu(5600000))

    # Slide 11: Thank You Slide
    s11 = prs.slides[10]
    for s in list(s11.shapes):
        sp = s._element
        sp.getparent().remove(sp)

    ty_box = s11.shapes.add_textbox(Emu(1000000), Emu(2600000), Emu(7900000), Emu(1200000))
    p_ty = ty_box.text_frame.paragraphs[0]
    p_ty.text = "감사합니다."
    p_ty.font.bold = True
    p_ty.font.size = Pt(40)
    p_ty.font.color.rgb = RGBColor(30, 41, 59)
    p_ty.alignment = PP_ALIGN.CENTER

    cp_box = s11.shapes.add_textbox(Emu(1000000), Emu(6300000), Emu(7900000), Emu(400000))
    p_cp = cp_box.text_frame.paragraphs[0]
    p_cp.text = "Corporate Introduction @ DANAL Co., Ltd. All Rights Reserved."
    p_cp.font.size = Pt(10)
    p_cp.font.color.rgb = RGBColor(148, 163, 184)
    p_cp.alignment = PP_ALIGN.CENTER

    out1 = os.path.join(APP_DIR, "신용카드 검수요청서(럭키픽).pptx")
    out2 = os.path.join(DOWNLOADS_DIR, "신용카드 검수요청서(럭키픽).pptx")
    prs.save(out1)
    prs.save(out2)
    print("Master 11-slide presentation generated successfully!")

# ===========================================================================
# 2. BUILD DEDICATED NEW 5-SLIDE PPT (Cover + 3 Focus Slides + Thank You)
# ===========================================================================
def build_dedicated_focus_slides():
    prs = pptx.Presentation(TEMPLATE_PATH)
    # Keep only 5 slides
    while len(prs.slides) > 5:
        delete_slide(prs, len(prs.slides) - 1)
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    while len(prs.slides) < 5:
        prs.slides.add_slide(blank_layout)

    # Slide 1: Cover
    s0 = prs.slides[0]
    for shape in s0.shapes:
        if shape.has_text_frame:
            text = shape.text.strip()
            tf = shape.text_frame
            if "업체명" in text or "다날" in text or "상호명" in text or shape.name == "텍스트 상자 28":
                tf.text = "럭키픽 (LuckyPick)\n포인트 사용 프로세스 및 환불 규정 검수요청서"
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].font.size = Pt(26)
                tf.paragraphs[0].font.color.rgb = RGBColor(26, 32, 44)
            elif "URL" in text or "TEST" in text or "작성" in text:
                tf.text = (
                    "서비스 URL: https://luckypick-ec4cf.web.app\n"
                    "상호: 럭키픽 | 대표: 이재영 | 사업자번호: 803-05-03449\n"
                    "고객센터: 070-8019-3205 | 이메일: leejeayoung0713@gmail.com\n"
                    "검수 항목: 8p 포인트 사용 프로세스 / 9p 이용약관(환불) / 10p 포인트 사용처\n"
                    "작성일자: 2026. 10. 21"
                )
                for p in tf.paragraphs:
                    p.font.size = Pt(13)
                    p.font.color.rgb = RGBColor(55, 65, 81)

    # Slide 2: 포인트 사용 프로세스 (Page 8)
    s1 = prs.slides[1]
    setup_header(
        s1,
        title_text="포인트 사용 프로세스",
        top_right_text="원하는 서비스 및 수량 선택 후\n결제를 진행하면 충전된 포인트에서 차감됩니다"
    )
    add_image(s1, r"C:\App\LuckyPick\danal_captures\08_point_use_process_v2.png", custom_top=Emu(950000), custom_max_h=Emu(5600000))

    # Slide 3: 이용약관(환불) (Page 9)
    s2 = prs.slides[2]
    setup_header(s2, title_text="이용약관(환불)")
    add_image(s2, r"C:\App\LuckyPick\danal_captures\09_refund_terms_v2.png", custom_top=Emu(950000), custom_max_h=Emu(5600000))

    # Slide 4: 포인트 사용처 (Page 10)
    s3 = prs.slides[3]
    setup_header(s3, title_text="포인트 사용처")
    add_image(s3, r"C:\App\LuckyPick\danal_captures\10_point_scope_v2.png", custom_top=Emu(950000), custom_max_h=Emu(5600000))

    # Slide 5: Thank You Slide
    s4 = prs.slides[4]
    for s in list(s4.shapes):
        sp = s._element
        sp.getparent().remove(sp)

    ty_box = s4.shapes.add_textbox(Emu(1000000), Emu(2600000), Emu(7900000), Emu(1200000))
    p_ty = ty_box.text_frame.paragraphs[0]
    p_ty.text = "감사합니다."
    p_ty.font.bold = True
    p_ty.font.size = Pt(40)
    p_ty.font.color.rgb = RGBColor(30, 41, 59)
    p_ty.alignment = PP_ALIGN.CENTER

    cp_box = s4.shapes.add_textbox(Emu(1000000), Emu(6300000), Emu(7900000), Emu(400000))
    p_cp = cp_box.text_frame.paragraphs[0]
    p_cp.text = "Corporate Introduction @ DANAL Co., Ltd. All Rights Reserved."
    p_cp.font.size = Pt(10)
    p_cp.font.color.rgb = RGBColor(148, 163, 184)
    p_cp.alignment = PP_ALIGN.CENTER

    out1 = os.path.join(APP_DIR, "포인트_및_환불_검수요청서(럭키픽).pptx")
    out2 = os.path.join(DOWNLOADS_DIR, "포인트_및_환불_검수요청서(럭키픽).pptx")
    prs.save(out1)
    prs.save(out2)
    print("Dedicated 5-slide presentation generated successfully!")

if __name__ == "__main__":
    build_master_11_slides()
    build_dedicated_focus_slides()
