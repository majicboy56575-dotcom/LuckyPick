import os
import sys
import io
import pptx
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from PIL import Image

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def build_danal_verification_pptx():
    template_path = r"C:\Users\USER\Downloads\신용카드 검수요청서(양식).pptx"
    out_path = r"C:\App\LuckyPick\신용카드 검수요청서(럭키픽).pptx"
    downloads_path = r"C:\Users\USER\Downloads\신용카드 검수요청서(럭키픽).pptx"

    prs = pptx.Presentation(template_path)
    print(f"Initial template slides count: {len(prs.slides)}")

    def delete_slide(prs, index):
        rId = prs.slides._sldIdLst[index].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[index]

    # Keep 11 slides (indices 0 to 10)
    while len(prs.slides) > 11:
        delete_slide(prs, len(prs.slides) - 1)

    # If less than 11, we can add slides using a blank layout
    blank_layout = prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1]
    while len(prs.slides) < 11:
        prs.slides.add_slide(blank_layout)

    print(f"Target slides count: {len(prs.slides)}")

    # -------------------------------------------------------------
    # Slide 1 (Index 0): Cover Slide
    # -------------------------------------------------------------
    s0 = prs.slides[0]
    for shape in s0.shapes:
        if shape.has_text_frame:
            text = shape.text.strip()
            tf = shape.text_frame
            if "업체명" in text or "다날" in text or "상호명" in text or shape.name == "텍스트 상자 28":
                tf.text = "럭키픽 (LuckyPick)\n다날 결제서비스 연동 검수요청서"
                tf.paragraphs[0].font.bold = True
                tf.paragraphs[0].font.size = Pt(28)
                tf.paragraphs[0].font.color.rgb = RGBColor(26, 32, 44)
            elif "URL" in text or "TEST" in text or "작성" in text:
                tf.text = (
                    "서비스 URL: https://luckypick-ec4cf.web.app\n"
                    "상호(법인명): 럭키픽 | 대표자: 이재영\n"
                    "사업자등록번호: 803-05-03449 | 통신판매업: 제 2026-화성새솔-0099 호\n"
                    "고객센터: 070-8019-3205 | 이메일: leejeayoung0713@gmail.com\n"
                    "사업장 소재지: 경기도 화성시 만세구 수노을1로 148, 101동 902호\n"
                    "TEST 계정: ID: tester@luckypick.kr / PW: LuckyPick2026!\n"
                    "신청 서비스: 휴대폰 소액결제(APP), 신용카드 결제(일반/간편)\n"
                    "작성일자: 2026. 10. 21"
                )
                for p in tf.paragraphs:
                    p.font.size = Pt(13)
                    p.font.color.rgb = RGBColor(55, 65, 81)

    # -------------------------------------------------------------
    # Slide Image Layout Helper
    # -------------------------------------------------------------
    max_w = Emu(9106000)
    max_h = Emu(5700000)
    base_top = Emu(850000)
    base_left = Emu(400000)

    def setup_slide_header(slide, title_text, subtitle_text=None, test_account=False):
        # Remove old body shapes
        shapes_to_remove = []
        title_shape = None
        for shape in slide.shapes:
            if shape.has_text_frame and shape.top < Emu(750000):
                title_shape = shape
            else:
                shapes_to_remove.append(shape)

        for s in shapes_to_remove:
            sp = s._element
            sp.getparent().remove(sp)

        if not title_shape:
            title_shape = slide.shapes.add_textbox(Emu(300000), Emu(120000), Emu(9300000), Emu(650000))

        tf = title_shape.text_frame
        tf.word_wrap = True
        tf.text = title_text
        p0 = tf.paragraphs[0]
        p0.font.bold = True
        p0.font.size = Pt(17)
        p0.font.color.rgb = RGBColor(30, 41, 59)

        if subtitle_text:
            p1 = tf.add_paragraph()
            p1.text = subtitle_text
            p1.font.bold = False
            p1.font.size = Pt(13)
            p1.font.color.rgb = RGBColor(71, 85, 105)

        if test_account:
            tb_test = slide.shapes.add_textbox(Emu(6800000), Emu(120000), Emu(2800000), Emu(500000))
            tf_t = tb_test.text_frame
            tf_t.text = "TEST ID : tester@luckypick.kr\nTEST PW : LuckyPick2026!"
            for p in tf_t.paragraphs:
                p.font.size = Pt(11)
                p.font.bold = True
                p.font.color.rgb = RGBColor(220, 38, 38)
                p.alignment = PP_ALIGN.RIGHT

    def add_image_to_slide(slide, img_path):
        if not os.path.exists(img_path):
            print(f"Warning: Image not found: {img_path}")
            return
        with Image.open(img_path) as im:
            w_px, h_px = im.size

        scale_w = max_w / w_px
        scale_h = max_h / h_px
        scale = min(scale_w, scale_h)

        final_w = int(w_px * scale)
        final_h = int(h_px * scale)

        pos_left = int(base_left + (max_w - final_w) / 2)
        pos_top = int(base_top + (max_h - final_h) / 2)

        slide.shapes.add_picture(img_path, pos_left, pos_top, final_w, final_h)

    # -------------------------------------------------------------
    # Slides 2 to 9 Data (Image Slides)
    # -------------------------------------------------------------
    slides_config = [
        # Slide 2 (Index 1): Main Home
        {
            "index": 1,
            "title": "① 사이트 메인 (URL노출 캡쳐)",
            "subtitle": "서비스 URL (https://luckypick-ec4cf.web.app) 상단 주소창 노출",
            "img": r"C:\App\LuckyPick\danal_captures\01_main_home.png",
            "test_account": False
        },
        # Slide 3 (Index 2): Footer Business Info
        {
            "index": 2,
            "title": "② 사이트 메인 하단정보 ※상호명, 사업자번호, 유선번호, 소재지, 대표자명 기재必",
            "subtitle": "상호: 럭키픽 | 대표: 이재영 | 사업자번호: 803-05-03449 | 고객센터: 070-8019-3205",
            "img": r"C:\App\LuckyPick\danal_captures\02_footer_business_info.png",
            "test_account": False
        },
        # Slide 4 (Index 3): Product & Point Package Selection
        {
            "index": 3,
            "title": "③ 상품 및 포인트 충전 결제 프로세스 ※ 결제진행상품금액동일/ 금액임의입력불가",
            "subtitle": "정찰제 패키지 선택 (10,000P / 32,000P / 55,000P / 115,000P), 금액 임의 수정 불가",
            "img": r"C:\App\LuckyPick\danal_captures\07_point_charge_section.png",
            "test_account": True
        },
        # Slide 5 (Index 4): Payment Selection Modal
        {
            "index": 4,
            "title": "④ 결제창 연동 (포인트 충전 ₩10,000원 결제 수단 선택)",
            "subtitle": "정찰제 ₩10,000원 결제 대상 > 신용·체크카드 / 휴대폰 소액결제 선택창 호출",
            "img": r"C:\App\LuckyPick\danal_captures\08_point_charge_modal.png",
            "test_account": True
        },
        # Slide 6 (Index 5): Danal Card Window
        {
            "index": 5,
            "title": "⑤ 각 카드사 인증창 호출 (다날 신용카드 표준결제창 호출)",
            "subtitle": "다날 신용카드 표준 결제창 호출 > 각 카드사 선택창 (10,000원 결제)",
            "img": r"C:\Users\USER\.gemini\antigravity-ide\brain\d7311755-2ddf-4acd-8b6e-a84b6e5f359a\.user_uploaded\media_1792501582882.png",
            "test_account": False
        },
        # Slide 7 (Index 6): Danal Phone Payment Window
        {
            "index": 6,
            "title": "⑤ 휴대폰 소액결제 연동창 호출 (SKT, KT, LGU+, 알뜰폰)",
            "subtitle": "다날 휴대폰 소액결제 연동창 호출 > 통신 3사 및 알뜰폰 본인인증 결제창",
            "img": r"C:\Users\USER\.gemini\antigravity-ide\brain\d7311755-2ddf-4acd-8b6e-a84b6e5f359a\.user_uploaded\media_1792501062466.png",
            "test_account": False
        },
        # Slide 8 (Index 7): Point Usage Process (Sample Page 8)
        {
            "index": 7,
            "title": "포인트 사용 프로세스",
            "subtitle": "원하는 럭키박스 선택 후 개봉을 진행하면 충전된 포인트에서 차감됩니다",
            "img": r"C:\App\LuckyPick\danal_captures\09_point_use_process.png",
            "test_account": False
        },
        # Slide 9 (Index 8): Refund Policy (Sample Page 9)
        {
            "index": 8,
            "title": "이용약관(환불)",
            "subtitle": "충전 후 미사용 포인트 7일 이내 원 결제수단 100% 전액 환불 규정",
            "img": r"C:\App\LuckyPick\danal_captures\10_point_refund_policy.png",
            "test_account": False
        }
    ]

    for cfg in slides_config:
        slide = prs.slides[cfg["index"]]
        setup_slide_header(slide, cfg["title"], cfg.get("subtitle"), cfg.get("test_account", False))
        add_image_to_slide(slide, cfg["img"])
        print(f"Slide {cfg['index']+1} generated: {cfg['title']}")

    # -------------------------------------------------------------
    # Slide 10 (Index 9): Point Usage Restriction (Sample Page 10)
    # -------------------------------------------------------------
    s10 = prs.slides[9]
    setup_slide_header(s10, "포인트 사용처", None, False)

    # Big red notice in center like sample page 10
    notice_box = s10.shapes.add_textbox(Emu(600000), Emu(1600000), Emu(8700000), Emu(1000000))
    tf_n = notice_box.text_frame
    tf_n.word_wrap = True
    p_n = tf_n.paragraphs[0]
    p_n.text = "해당 사이트에서만 포인트 사용이 가능합니다."
    p_n.font.bold = True
    p_n.font.size = Pt(26)
    p_n.font.color.rgb = RGBColor(220, 38, 38) # Strong Red
    p_n.alignment = PP_ALIGN.LEFT

    # Detailed Explanation Card Box
    desc_box = s10.shapes.add_textbox(Emu(600000), Emu(2700000), Emu(8700000), Emu(3600000))
    tf_d = desc_box.text_frame
    tf_d.word_wrap = True

    items = [
        ("1. 포인트 사용 가능 사이트", "• 공식 서비스 URL: https://luckypick-ec4cf.web.app (럭키픽 온라인 플랫폼 전용)\n• 당사 사이트 외 타 사이트, 오프라인 매장 등에서는 일체 사용이 불가능합니다."),
        ("2. 포인트 사용 용도", "• 럭키픽 플랫폼 내 블라인드 큐레이션 럭키박스(베이직, 프리미엄, VIP) 개봉 이용\n• 실물 득템 상품 보관함 출고 시 택배 배송비(3,000P) 결제 이용"),
        ("3. 환금성 및 외부 유출 원천 차단 (비환금성 준수)", "• 현금 환전 및 계좌 인출이 절대 불가능합니다. (환금성 0%)\n• 회원 간 포인트 선물, 양도, 이전 및 장외 개인 간 거래가 시스템상 원천 차단되어 있습니다."),
        ("4. 미사용 포인트 청약철회 및 환불 원칙", "• 유료 결제로 충전한 후 사용하지 않은 잔여 포인트는 결제일로부터 7일 이내 요청 시 구매자의 원 결제수단(신용카드 승인 취소, 휴대폰 소액결제 승인 취소)으로 100% 전액 환불 처리됩니다.\n• 별도의 현금 지급이나 계좌 송금 형태의 환불은 전자상거래법 및 PG사 규정에 따라 불가합니다.")
    ]

    for title, content in items:
        p_t = tf_d.add_paragraph() if tf_d.paragraphs[0].text else tf_d.paragraphs[0]
        p_t.text = title
        p_t.font.bold = True
        p_t.font.size = Pt(14)
        p_t.font.color.rgb = RGBColor(15, 23, 42)
        p_t.space_before = Pt(8)

        p_c = tf_d.add_paragraph()
        p_c.text = content
        p_c.font.size = Pt(12)
        p_c.font.color.rgb = RGBColor(71, 85, 105)
        p_c.space_after = Pt(6)

    print("Slide 10 generated: 포인트 사용처")

    # -------------------------------------------------------------
    # Slide 11 (Index 10): Thank You Slide (Sample Page 11)
    # -------------------------------------------------------------
    s11 = prs.slides[10]
    # Remove all old shapes
    for s in list(s11.shapes):
        sp = s._element
        sp.getparent().remove(sp)

    # Add thank you centered
    ty_box = s11.shapes.add_textbox(Emu(1000000), Emu(2600000), Emu(7900000), Emu(1200000))
    tf_ty = ty_box.text_frame
    p_ty = tf_ty.paragraphs[0]
    p_ty.text = "감사합니다."
    p_ty.font.bold = True
    p_ty.font.size = Pt(40)
    p_ty.font.color.rgb = RGBColor(30, 41, 59)
    p_ty.alignment = PP_ALIGN.CENTER

    # Add footer copyright
    cp_box = s11.shapes.add_textbox(Emu(1000000), Emu(6300000), Emu(7900000), Emu(400000))
    tf_cp = cp_box.text_frame
    p_cp = tf_cp.paragraphs[0]
    p_cp.text = "Corporate Introduction @ DANAL Co., Ltd. All Rights Reserved."
    p_cp.font.size = Pt(10)
    p_cp.font.color.rgb = RGBColor(148, 163, 184)
    p_cp.alignment = PP_ALIGN.CENTER

    print("Slide 11 generated: 감사합니다.")

    # -------------------------------------------------------------
    # Save Presentation
    # -------------------------------------------------------------
    prs.save(out_path)
    prs.save(downloads_path)
    print(f"\nSUCCESS! 11-slide PPTX generated successfully:")
    print(f"1) {out_path}")
    print(f"2) {downloads_path}")

if __name__ == "__main__":
    build_danal_verification_pptx()
